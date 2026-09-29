"""Выгрузка результатов в книгу Excel для независимой проверки расчёта.

В книгу записываются те же псевдослучайные числа, которые использовала программа,
а затраты пересчитываются формулами Excel. Если формулы дают те же суммы, что и
программа, значит расчёт выполнен правильно. Excel считает значения при открытии файла.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import Parameters, SimulationResult, run_percent

TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
GOOD_FILL = PatternFill("solid", fgColor="E2EFDA")
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
WHITE_BOLD = Font(color="FFFFFF", bold=True, size=12)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="B7C7DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _title(sheet, text: str, width: int) -> None:
    """Заголовок листа на всю ширину таблицы."""
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = sheet.cell(row=1, column=1, value=text)
    cell.fill = TITLE_FILL
    cell.font = WHITE_BOLD
    cell.alignment = CENTER
    sheet.row_dimensions[1].height = 26


def _header(sheet, row: int, values: list[str]) -> None:
    """Строка заголовков таблицы."""
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.fill = HEAD_FILL
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = BORDER


def build_workbook(result: SimulationResult, excel_trials: int = 2000) -> Workbook:
    """Сборка книги Excel по результатам моделирования."""
    params = result.parameters
    count = min(excel_trials, params.trials)
    numbers = result.numbers[:count]
    percents = params.percents()

    book = Workbook()
    _sheet_parameters(book.active, params, count)
    _sheet_trials(book.create_sheet("Испытания"), params, numbers, percents)
    _sheet_summary(book.create_sheet("Свод"), params, result, numbers, percents, count)
    _sheet_example(book.create_sheet("Пример испытаний"), result)
    return book


def _sheet_parameters(sheet, params: Parameters, count: int) -> None:
    """Лист с исходными данными. На него ссылаются формулы остальных листов."""
    sheet.title = "Параметры"
    _title(sheet, "Исходные данные задачи (вариант 3)", 3)
    rows = [
        ("A — доля негодных резисторов, %", params.defect_percent, "разыгрывается числом R2"),
        ("B — стоимость контроля одного резистора, руб.", params.cost_control, "платится за каждый проверенный резистор"),
        ("C — стоимость подгонки, руб.", params.cost_fitting, "если брак выявлен при контроле"),
        ("D — стоимость замены в изделии, руб.", params.cost_replace, "если брак не выявлен"),
        ("Испытаний в программе", params.trials, "количество резисторов"),
        ("Испытаний в этом файле", count, "столько строк выгружено для проверки"),
        ("Шаг перебора доли контроля, %", params.percent_step, ""),
        ("Номер полинома LFSR", params.polynomial, params.polynomial_text()),
        ("Начальное состояние регистра", params.seed, "seed генератора"),
    ]
    _header(sheet, 2, ["Показатель", "Значение", "Пояснение"])
    for index, (name, value, note) in enumerate(rows, start=3):
        sheet.cell(row=index, column=1, value=name).border = BORDER
        cell = sheet.cell(row=index, column=2, value=value)
        cell.border = BORDER
        cell.font = BOLD
        cell.alignment = CENTER
        sheet.cell(row=index, column=3, value=note).border = BORDER
    sheet.column_dimensions["A"].width = 44
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 60
    # Имена ячеек делают формулы на других листах понятными.
    book = sheet.parent
    book.defined_names.add(_defined_name("Брак", "Параметры", "$B$3"))
    book.defined_names.add(_defined_name("ЗатрКонтроль", "Параметры", "$B$4"))
    book.defined_names.add(_defined_name("ЗатрПодгонка", "Параметры", "$B$5"))
    book.defined_names.add(_defined_name("ЗатрЗамена", "Параметры", "$B$6"))


def _defined_name(name: str, sheet: str, ref: str):
    from openpyxl.workbook.defined_name import DefinedName

    return DefinedName(name, attr_text=f"'{sheet}'!{ref}")


def _sheet_trials(sheet, params: Parameters, numbers, percents) -> None:
    """Лист с псевдослучайными числами и формулами затрат по каждой доле контроля."""
    _title(sheet, "Пересчёт затрат формулами Excel по тем же псевдослучайным числам", 3 + len(percents))
    sheet.cell(row=2, column=1, value="Доля контроля, %").font = BOLD
    for column, percent in enumerate(percents, start=4):
        cell = sheet.cell(row=2, column=column, value=percent / 100)
        cell.number_format = "0%"
        cell.font = BOLD
        cell.alignment = CENTER
        cell.fill = HEAD_FILL
        cell.border = BORDER
    _header(sheet, 3, ["№", "R1", "R2"] + ["затраты" for _ in percents])

    for index, (r1, r2) in enumerate(numbers, start=1):
        row = 3 + index
        sheet.cell(row=row, column=1, value=index)
        sheet.cell(row=row, column=2, value=r1).number_format = "0.000000"
        sheet.cell(row=row, column=3, value=r2).number_format = "0.000000"
        for column, _ in enumerate(percents, start=4):
            letter = get_column_letter(column)
            formula = (
                f"=IF($B{row}<{letter}$2,ЗатрКонтроль,0)"
                f"+IF(AND($B{row}<{letter}$2,$C{row}<Брак/100),ЗатрПодгонка,0)"
                f"+IF(AND($B{row}>={letter}$2,$C{row}<Брак/100),ЗатрЗамена,0)"
            )
            sheet.cell(row=row, column=column, value=formula).number_format = "0"

    total_row = 4 + len(numbers)
    sheet.cell(row=total_row, column=1, value="Итого").font = BOLD
    for column, _ in enumerate(percents, start=4):
        letter = get_column_letter(column)
        cell = sheet.cell(row=total_row, column=column, value=f"=SUM({letter}4:{letter}{total_row - 1})")
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.number_format = "#,##0"
    sheet.freeze_panes = "D4"
    sheet.column_dimensions["A"].width = 6
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12
    for column in range(4, 4 + len(percents)):
        sheet.column_dimensions[get_column_letter(column)].width = 10


def _sheet_summary(sheet, params: Parameters, result: SimulationResult, numbers, percents, count: int) -> None:
    """Главный лист проверки: сравнение итогов программы и формул Excel."""
    _title(sheet, "Сравнение результатов программы и формул Excel", 6)
    _header(sheet, 2, [
        "Доля контроля, %",
        "Затраты: программа",
        "Затраты: формулы Excel",
        "Разница",
        "Затраты: теоретическая формула",
        "Отклонение модели от теории, %",
    ])

    # Программа пересчитывает свои итоги по тем же строкам, что выгружены в файл.
    control_rows = [run_percent(params, percent, numbers) for percent in percents]
    trials_total_row = 4 + len(numbers)

    for index, (percent, row_result) in enumerate(zip(percents, control_rows), start=3):
        letter = get_column_letter(3 + index - 2)  # столбец листа «Испытания» для этой доли
        sheet.cell(row=index, column=1, value=percent).alignment = CENTER
        sheet.cell(row=index, column=2, value=row_result.total_cost).number_format = "#,##0"
        sheet.cell(row=index, column=3, value=f"=Испытания!{letter}{trials_total_row}").number_format = "#,##0"
        diff = sheet.cell(row=index, column=4, value=f"=B{index}-C{index}")
        diff.number_format = "#,##0"
        share = percent / 100
        sheet.cell(
            row=index,
            column=5,
            value=f"={count}*({share}*ЗатрКонтроль+{share}*Брак/100*ЗатрПодгонка+(1-{share})*Брак/100*ЗатрЗамена)",
        ).number_format = "#,##0"
        sheet.cell(row=index, column=6, value=f"=(B{index}-E{index})/E{index}*100").number_format = "0.00"
        for column in range(1, 7):
            sheet.cell(row=index, column=column).border = BORDER

    last = 2 + len(percents)
    check_row = last + 2
    sheet.cell(row=check_row, column=1, value="Проверка совпадения").font = BOLD
    sheet.cell(
        row=check_row,
        column=2,
        value=f'=IF(SUMPRODUCT(--(ABS(B3:B{last}-C3:C{last})>0.001))=0,"Совпадает: формулы Excel дают те же затраты","Есть расхождения")',
    ).font = BOLD
    sheet.cell(row=check_row, column=2).fill = GOOD_FILL

    best_row = check_row + 1
    sheet.cell(row=best_row, column=1, value="Лучшая доля контроля по Excel, %").font = BOLD
    sheet.cell(row=best_row, column=2, value=f"=INDEX(A3:A{last},MATCH(MIN(C3:C{last}),C3:C{last},0))").font = BOLD
    sheet.cell(row=best_row, column=2).fill = GOOD_FILL
    sheet.cell(row=best_row + 1, column=1, value="Лучшая доля контроля по программе, %").font = BOLD
    sheet.cell(row=best_row + 1, column=2, value=min(control_rows, key=lambda r: r.total_cost).percent).font = BOLD
    sheet.cell(row=best_row + 1, column=2).fill = GOOD_FILL
    sheet.cell(
        row=best_row + 3,
        column=1,
        value="Столбец «Затраты: формулы Excel» считается на листе «Испытания» по тем же числам R1 и R2, "
              "которые использовала программа. Совпадение столбцов B и C подтверждает правильность расчёта.",
    )
    sheet.merge_cells(start_row=best_row + 3, start_column=1, end_row=best_row + 3, end_column=6)
    sheet.cell(row=best_row + 3, column=1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[best_row + 3].height = 34

    chart = LineChart()
    chart.title = "Затраты в зависимости от доли контроля"
    chart.y_axis.title = "Затраты, руб."
    chart.x_axis.title = "Доля контроля, %"
    data = Reference(sheet, min_col=2, max_col=3, min_row=2, max_row=last)
    cats = Reference(sheet, min_col=1, min_row=3, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 9, 18
    sheet.add_chart(chart, f"H3")

    sheet.column_dimensions["A"].width = 34
    for column in "BCDEF":
        sheet.column_dimensions[column].width = 22


def _sheet_example(sheet, result: SimulationResult) -> None:
    """Лист с подробной таблицей испытаний — как таблица 1 в методических указаниях."""
    _title(sheet, f"Пример испытаний при доле контроля {result.example_percent} %", 7)
    _header(sheet, 2, ["Номер испытания", "R1", "Контроль", "R2", "Требуется подгонка", "Затраты", "Общие затраты"])
    for index, trial in enumerate(result.trials_example, start=3):
        sheet.cell(row=index, column=1, value=trial.number).alignment = CENTER
        sheet.cell(row=index, column=2, value=round(trial.r1, 4)).number_format = "0.0000"
        sheet.cell(row=index, column=3, value="да" if trial.controlled else "нет").alignment = CENTER
        sheet.cell(row=index, column=4, value=round(trial.r2, 4)).number_format = "0.0000"
        sheet.cell(row=index, column=5, value="да" if trial.defective else "нет").alignment = CENTER
        sheet.cell(row=index, column=6, value=trial.cost).number_format = "0"
        sheet.cell(row=index, column=7, value=trial.total).number_format = "#,##0"
        if trial.cost:
            for column in range(1, 8):
                sheet.cell(row=index, column=column).fill = WARN_FILL
        for column in range(1, 8):
            sheet.cell(row=index, column=column).border = BORDER
    for column, width in zip("ABCDEFG", (18, 12, 12, 12, 20, 12, 16)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"


def save_workbook(result: SimulationResult, path: str, excel_trials: int = 2000) -> str:
    """Сохранение книги Excel на диск. Возвращает путь к файлу."""
    book = build_workbook(result, excel_trials)
    book.save(path)
    return path
